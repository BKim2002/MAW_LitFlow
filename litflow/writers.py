from __future__ import annotations

from pathlib import Path
from typing import Any
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .agents import format_citation
from .models import Record, RunConfig
from .notion_writer import NotionPackagingAgent
from .utils import ensure_dir, read_json, read_jsonl


class PackagingAgent:
    name = "PackagingAgent"

    def __init__(self, notion_client: Any | None = None) -> None:
        self.notion_client = notion_client

    def run(self, records: list[Record], summaries: list[dict[str, Any]], flags: list[dict[str, Any]], config: RunConfig, out_dir: Path) -> dict[str, str]:
        outputs: dict[str, str] = {}
        if config.output_format not in {"files", "notion", "both"}:
            raise RuntimeError(f"Unsupported output_format: {config.output_format}")
        if config.output_format in {"files", "both"}:
            xlsx_path = out_dir / "literature_review.xlsx"
            docx_path = out_dir / "literature_review.docx"
            write_xlsx(xlsx_path, records, summaries, flags, config, out_dir)
            write_docx(docx_path, records, summaries, flags, config, out_dir)
            audit_docx_structure(docx_path)
            outputs.update({"xlsx": str(xlsx_path), "docx": str(docx_path)})
        if config.output_format in {"notion", "both"}:
            outputs.update(NotionPackagingAgent(self.notion_client).run(records, summaries, flags, config, out_dir))
        return outputs


def write_xlsx(path: Path, records: list[Record], summaries: list[dict[str, Any]], flags: list[dict[str, Any]], config: RunConfig, out_dir: Path) -> None:
    ensure_dir(path.parent)
    wb = Workbook()
    wb.remove(wb.active)
    add_sheet(wb, "Run_Config", [{"key": k, "value": v} for k, v in config.to_dict().items()])
    add_sheet(wb, "Search_Log", read_jsonl(out_dir / "search_log.jsonl"))
    query_plan = read_json(out_dir / "query_plan.json") if (out_dir / "query_plan.json").exists() else {"queries": []}
    add_sheet(wb, "Queries", query_plan.get("queries", []))
    record_rows = [record_to_row(r) for r in records]
    add_sheet(wb, "All_Candidates", record_rows)
    add_sheet(wb, "Included", [record_to_row(r) for r in records if r.inclusion_status == "included"])
    add_sheet(wb, "Excluded", [record_to_row(r) for r in records if r.inclusion_status == "excluded"])
    add_sheet(wb, "Evidence_Level", [{"record_id": r.record_id, "title": r.title, "evidence_level": r.evidence_level, "open_access_pdf": r.open_access_pdf} for r in records])
    add_sheet(wb, "Deep_Summaries", summaries)
    add_sheet(wb, "FullText_Requests", read_jsonl(out_dir / "fulltext_requests.jsonl"))
    add_sheet(wb, "QA_Flags", flags)
    wb.save(path)


def add_sheet(wb: Workbook, name: str, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(title=name[:31])
    if not rows:
        ws.append(["status"])
        ws.append(["no rows"])
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([normalize_cell(row.get(header)) for header in headers])
    header_fill = PatternFill("solid", fgColor="E8EEF5")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for idx, header in enumerate(headers, start=1):
        width = min(max(len(str(header)) + 2, 12), 45)
        for cell in ws[get_column_letter(idx)][1:30]:
            width = min(max(width, min(len(str(cell.value or "")) + 2, 45)), 60)
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"


def normalize_cell(value: Any) -> Any:
    if isinstance(value, (list, dict)):
        import json

        return json.dumps(value, ensure_ascii=False)
    return value


def record_to_row(record: Record) -> dict[str, Any]:
    return {
        "record_id": record.record_id,
        "title": record.title,
        "authors": "; ".join(record.authors),
        "year": record.year,
        "venue": record.venue,
        "doi": record.doi,
        "source_ids": record.source_ids,
        "abstract": record.abstract,
        "url": record.url,
        "source_database": "; ".join(record.source_database),
        "evidence_level": record.evidence_level,
        "relevance_score": record.relevance_score,
        "inclusion_status": record.inclusion_status,
        "exclusion_reason": record.exclusion_reason,
        "summary_status": record.summary_status,
        "citation_count": record.citation_count,
        "open_access_pdf": record.open_access_pdf,
        "found_by": record.found_by,
        "search_round": record.search_round,
        "facet_matches": record.facet_matches,
        "coverage_warning": record.coverage_warning,
    }


def write_docx(path: Path, records: list[Record], summaries: list[dict[str, Any]], flags: list[dict[str, Any]], config: RunConfig, out_dir: Path) -> None:
    ensure_dir(path.parent)
    doc = Document()
    configure_docx_styles(doc)
    add_title(doc, "문헌검색 결과 요약")
    add_meta(doc, config, records)
    doc.add_heading("조사 범위와 방법", level=1)
    add_para(doc, f"주제: {config.topic}")
    add_para(doc, "이 문서는 litflow 로컬 멀티 에이전트 워크플로우가 생성한 자동 문헌검색 결과입니다.")
    add_para(doc, "공개 커넥터는 OpenAlex, Crossref, Semantic Scholar, arXiv를 사용하며, Google Scholar/Scopus/Web of Science 등은 별도 검색식 패키지로 남깁니다.")
    doc.add_heading("검색식 요약", level=1)
    query_plan = read_json(out_dir / "query_plan.json") if (out_dir / "query_plan.json").exists() else {"queries": []}
    for query in query_plan.get("queries", [])[:10]:
        if query.get("source") == "openalex":
            add_labeled_para(doc, query.get("intent", "query"), query.get("query", ""))
    doc.add_heading("커버리지 한계", level=1)
    add_para(doc, "v1은 유료 DB 내부 검색을 직접 보장하지 않습니다. 검색 로그와 manual_db_search_pack.md를 함께 검토하세요.")
    add_para(doc, "DOCX에는 URL을 넣지 않습니다. URL/DOI/source id는 XLSX에서 확인하세요.")
    doc.add_heading("핵심 문헌 맵", level=1)
    add_core_table(doc, records)
    doc.add_heading("문헌별 장문 요약", level=1)
    if summaries:
        for idx, summary in enumerate(summaries, start=1):
            doc.add_heading(f"{idx}. {summary.get('citation', summary.get('title', 'Untitled'))}", level=2)
            for key, label in [
                ("evidence_level", "근거 수준"),
                ("abstract", "Abstract 중심"),
                ("introduction", "Introduction 중심"),
                ("method", "Method 중심"),
                ("results_findings", "Results/Findings 중심"),
                ("conclusion", "Conclusion 중심"),
                ("limitations", "Limitations"),
                ("topic_relevance", "주제 관련성"),
                ("follow_up", "후속 검토"),
            ]:
                add_labeled_para(doc, label, str(summary.get(key, "")))
    else:
        add_para(doc, "포함 기준을 통과한 문헌 요약이 없습니다.")
    doc.add_heading("추가 DB 검색 필요 영역", level=1)
    add_para(doc, "Scopus, Web of Science, Google Scholar, ACM Digital Library, IEEE Xplore, Dimensions에서 manual_db_search_pack.md의 검색식을 재실행하고 export 파일을 추후 병합하세요.")
    doc.add_heading("QA 플래그", level=1)
    if flags:
        for item in flags:
            add_labeled_para(doc, item.get("code", "flag"), item.get("message", ""))
    else:
        add_para(doc, "QA 플래그가 없습니다.")
    doc.save(path)


def configure_docx_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25


def add_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(text)
    set_run(run, size=20, bold=True, color="0B2545")


def add_meta(doc: Document, config: RunConfig, records: list[Record]) -> None:
    included = len([r for r in records if r.inclusion_status == "included"])
    add_para(doc, f"출력 언어: {config.language} | 전체 후보: {len(records)} | 포함 문헌: {included}")


def add_para(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_run(run)


def add_labeled_para(doc: Document, label: str, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(6)
    r1 = p.add_run(f"{label}: ")
    set_run(r1, bold=True, color="1F4D78")
    r2 = p.add_run(text)
    set_run(r2)


def set_run(run, size: int | None = None, bold: bool = False, color: str | None = None) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.bold = bold
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_core_table(doc: Document, records: list[Record]) -> None:
    included = sorted([r for r in records if r.inclusion_status == "included"], key=lambda r: (r.relevance_score, r.citation_count or 0), reverse=True)[:20]
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    headers = ["문헌", "연도", "출처", "근거", "점수"]
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = header
        shade_cell(cell, "E8EEF5")
    for record in included:
        row = table.add_row()
        values = [
            format_citation(record),
            str(record.year or ""),
            "; ".join(record.source_database),
            record.evidence_level,
            str(record.relevance_score),
        ]
        for idx, value in enumerate(values):
            row.cells[idx].text = value
    set_table_geometry(table, [4200, 720, 1650, 1650, 1140])


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            for para in cell.paragraphs:
                para.paragraph_format.line_spacing = 1.15
                para.paragraph_format.space_after = Pt(0)
                if idx in {1, 4}:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    set_run(run, size=9 if idx == 0 else 10, bold=row is table.rows[0])


def set_cell_width(cell, width: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def audit_docx_structure(path: Path) -> None:
    with ZipFile(path) as zf:
        document_xml = zf.read("word/document.xml").decode("utf-8", errors="ignore")
        styles_xml = zf.read("word/styles.xml").decode("utf-8", errors="ignore")
    visible_text = extract_visible_docx_text(document_xml)
    required = {
        "tblGrid": "<w:tblGrid>" in document_xml,
        "table_widths": "w:tcW" in document_xml,
        "east_asia_font": "Malgun Gothic" in document_xml or "Malgun Gothic" in styles_xml,
        "no_docx_urls": "https://" not in visible_text and "http://" not in visible_text,
    }
    missing = [name for name, ok in required.items() if not ok]
    if missing:
        raise RuntimeError(f"DOCX structural audit failed: {missing}")


def extract_visible_docx_text(document_xml: str) -> str:
    try:
        root = ET.fromstring(document_xml)
    except ET.ParseError:
        return document_xml
    texts = []
    for node in root.iter():
        if node.tag.endswith("}t") and node.text:
            texts.append(node.text)
    return "\n".join(texts)
