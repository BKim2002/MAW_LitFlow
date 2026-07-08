from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import load_workbook

from litflow.agents import (
    CoverageAuditAgent,
    DeepSummaryAgent,
    EvidenceAcquisitionAgent,
    FullTextRequestAgent,
    IntakeScopeAgent,
    MetadataNormalizeDedupAgent,
    QualityAuditAgent,
    QueryStrategyAgent,
    RelevanceScreeningAgent,
    extract_elsevier_pii,
)
from litflow.connectors import BaseConnector, SearchResult, SerpApiGoogleScholarConnector
from litflow.models import RawRecord, Record, RunConfig, SearchLogEntry
from litflow.notion_writer import NotionPackagingAgent, hub_page_blocks, parse_notion_page_id, rich_text, summary_record_to_blocks
from litflow.orchestrator import Orchestrator
from litflow.sdk_bridge import LiteratureSummaryOutput, QABatchOutput, QAFlagOutput, QueryStrategyOutput, RelevanceBatchOutput, RelevanceDecision
from litflow.utils import read_json, read_jsonl, utc_now, write_json, write_jsonl
from litflow.writers import PackagingAgent


class FakeConnector(BaseConnector):
    def __init__(self, source: str, records: list[RawRecord]):
        self.source = source
        self.records = records

    def search(self, query: str, config: RunConfig, limit: int) -> SearchResult:
        rows = [RawRecord(**{**r.to_dict(), "source": self.source, "source_database": self.source}) for r in self.records[:limit]]
        return SearchResult(rows, SearchLogEntry(self.source, query, utc_now(), f"fake://{self.source}", len(rows)))


class ExplodingConnector(BaseConnector):
    source = "explode"

    def search(self, query: str, config: RunConfig, limit: int) -> SearchResult:
        raise AssertionError("resume-fulltext must not call search connectors")


class FakeSDKBridge:
    enabled = True

    def refine_query_strategy(self, topic, deterministic_plan):
        return QueryStrategyOutput(
            expanded_terms=["personnel selection"],
            concept_groups={"task_terms": ["employee selection"]},
            base_queries=["LLM employee selection validity"],
            rationale="fake sdk query strategy",
        )

    def screen_relevance(self, topic, records):
        return RelevanceBatchOutput(
            decisions=[
                RelevanceDecision(record_id=record.record_id, relevance_score=3, inclusion_status="included", exclusion_reason="")
                for record in records
            ]
        )

    def summarize_record(self, topic, record, evidence_bundle, fallback_citation):
        return LiteratureSummaryOutput(
            record_id=record.record_id,
            citation=fallback_citation,
            title=record.title,
            evidence_level=record.evidence_level,
            abstract="SDK abstract summary",
            introduction="SDK introduction summary",
            method="SDK method summary limited to evidence",
            results_findings="SDK findings summary limited to evidence",
            conclusion="SDK conclusion summary",
            limitations="SDK limitations",
            topic_relevance="SDK relevance",
            follow_up="SDK follow up",
        )

    def audit_quality(self, topic, records, summaries, existing_flags):
        return QABatchOutput(flags=[QAFlagOutput(record_id=records[0].record_id, title=records[0].title, code="sdk_flag", message="fake sdk flag")])


class FakeNotionClient:
    def __init__(self):
        self.pages_data = {}
        self.children = {}
        self.page_counter = 0
        self.block_counter = 0
        self.created_pages = []
        self.updated_pages = []
        self.deleted_blocks = []
        self.pages = FakeNotionPages(self)
        self.blocks = FakeNotionBlocks(self)

    def next_page_id(self):
        self.page_counter += 1
        return f"page{self.page_counter:032d}"[-32:]

    def next_block_id(self):
        self.block_counter += 1
        return f"block{self.block_counter:032d}"

    def page_title(self, properties):
        rich = properties["title"][0]["text"]["content"]
        return rich


class FakeNotionPages:
    def __init__(self, client):
        self.client = client

    def create(self, parent, properties):
        page_id = self.client.next_page_id()
        title = self.client.page_title(properties)
        page = {"id": page_id, "url": f"https://notion.test/{page_id}", "properties": properties, "title": title}
        self.client.pages_data[page_id] = page
        self.client.children.setdefault(page_id, [])
        self.client.created_pages.append(page_id)
        parent_id = parent.get("page_id")
        if parent_id:
            block_id = self.client.next_block_id()
            self.client.children.setdefault(parent_id, []).append({"id": block_id, "type": "child_page", "child_page": {"title": title}, "page_id": page_id})
        return page

    def update(self, page_id, properties):
        self.client.pages_data.setdefault(page_id, {"id": page_id, "url": f"https://notion.test/{page_id}"})
        self.client.pages_data[page_id]["properties"] = properties
        self.client.pages_data[page_id]["title"] = self.client.page_title(properties)
        self.client.updated_pages.append(page_id)
        return self.client.pages_data[page_id]


class FakeNotionChildren:
    def __init__(self, client):
        self.client = client

    def list(self, block_id, page_size=100, start_cursor=None):
        rows = self.client.children.get(block_id, [])
        start = int(start_cursor or 0)
        page = rows[start : start + page_size]
        next_cursor = start + page_size if start + page_size < len(rows) else None
        return {"results": page, "has_more": next_cursor is not None, "next_cursor": str(next_cursor) if next_cursor is not None else None}

    def append(self, block_id, children):
        stored = self.client.children.setdefault(block_id, [])
        for block in children:
            row = dict(block)
            row.setdefault("id", self.client.next_block_id())
            stored.append(row)
        return {"results": children}


class FakeNotionBlocks:
    def __init__(self, client):
        self.client = client
        self.children = FakeNotionChildren(client)

    def delete(self, block_id):
        for rows in self.client.children.values():
            for idx, row in enumerate(list(rows)):
                if row.get("id") == block_id:
                    rows.pop(idx)
                    self.client.deleted_blocks.append(block_id)
                    return {"id": block_id, "archived": True}
        self.client.deleted_blocks.append(block_id)
        return {"id": block_id, "archived": True}


class LitflowTests(unittest.TestCase):
    def test_korean_topic_expands_to_english_terms(self):
        with tempfile.TemporaryDirectory() as tmp:
            config = RunConfig(topic="LLM을 활용한 면접 평가와 역량 측정 연구", out=tmp)
            scope = IntakeScopeAgent().run(config, Path(tmp))
            expanded = " ".join(scope["expanded_terms"])
            self.assertIn("interview", expanded)
            self.assertIn("assessment", expanded)
            self.assertIn("competency", expanded)

    def test_high_recall_korean_topic_generates_faceted_queries(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp)
            config = RunConfig(topic="생성형 AI 활용 역량 혹은 AI Literacy와 직무성과와의 관계를 다룬 연구", out=tmp, recall_mode="high")
            scope = IntakeScopeAgent().run(config, out)
            plan = QueryStrategyAgent().run(scope, config, out)
            groups = plan["concept_groups"]
            self.assertIn("AI literacy", groups["core_concepts"])
            self.assertIn("job performance", groups["outcome_terms"])
            self.assertIn("workplace", groups["context_terms"])
            queries = plan["queries"]
            self.assertTrue(any(q["intent"] == "core_outcome" for q in queries))
            self.assertTrue(any(q["source"] == "serpapi_google_scholar" and q["search_round"] == 2 for q in queries))
            self.assertTrue(any("job performance" in q["query"] for q in queries))

    def test_coverage_audit_generates_supplemental_queries_for_weak_facets(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp)
            config = RunConfig(topic="AI literacy job performance", out=tmp, recall_mode="high")
            scope = IntakeScopeAgent().run(config, out)
            plan = QueryStrategyAgent().run(scope, config, out)
            record = Record(
                record_id="r1",
                title="AI Literacy Scale Development",
                abstract="This paper validates an AI literacy scale.",
                relevance_score=3,
                inclusion_status="included",
            )
            report = CoverageAuditAgent().run([record], plan, config, out, allow_supplemental=True)
            self.assertTrue(report["warnings"])
            self.assertTrue(report["supplemental_queries"])
            self.assertIn("outcome_terms", report["supplemental_queries"][0]["facets"])
            screened = read_jsonl(out / "screened_records.jsonl")[0]
            self.assertIn("facet_matches", screened)
            self.assertIn("coverage_warning", screened)

    def test_serpapi_missing_token_logs_warning(self):
        with tempfile.TemporaryDirectory() as tmp:
            config = RunConfig(topic="AI literacy", out=tmp, web_search_token_env="LITFLOW_MISSING_SERPAPI_TOKEN")
            result = SerpApiGoogleScholarConnector().search("AI literacy job performance", config, 5)
            self.assertEqual(result.records, [])
            self.assertIn("LITFLOW_MISSING_SERPAPI_TOKEN is not set", result.log.error)

    def test_serpapi_fixture_normalizes_google_scholar_result(self):
        fixture = {
            "organic_results": [
                {
                    "result_id": "scholar-1",
                    "title": "Generative AI Literacy and Job Performance",
                    "link": "https://example.test/paper",
                    "snippet": "This study examines generative AI literacy and employee job performance.",
                    "publication_info": {
                        "summary": "Kim, Lee - Journal of Work, 2025",
                        "authors": [{"name": "Kim"}, {"name": "Lee"}],
                    },
                    "inline_links": {"cited_by": {"total": 12}},
                    "resources": [{"file_format": "PDF", "link": "https://example.test/paper.pdf"}],
                }
            ]
        }
        with tempfile.TemporaryDirectory() as tmp:
            config = RunConfig(topic="AI literacy", out=tmp, web_search_token_env="SERPAPI_TEST_KEY")
            with patch.dict("os.environ", {"SERPAPI_TEST_KEY": "test-key"}), patch.object(SerpApiGoogleScholarConnector, "_get_json", return_value=fixture):
                result = SerpApiGoogleScholarConnector().search("AI literacy job performance", config, 5)
            self.assertEqual(len(result.records), 1)
            record = result.records[0]
            self.assertEqual(record.source, "serpapi_google_scholar")
            self.assertEqual(record.year, 2025)
            self.assertEqual(record.citation_count, 12)
            self.assertEqual(record.open_access_pdf, "https://example.test/paper.pdf")

    def test_dedup_merges_same_doi(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp)
            raw_dir = out / "raw_results"
            raw_dir.mkdir()
            write_jsonl(
                raw_dir / "a.jsonl",
                [
                    RawRecord(source="openalex", source_id="oa1", title="A Study of AI Interview Assessment", doi="10.123/ABC", abstract="First abstract", source_database="openalex").to_dict(),
                    RawRecord(source="crossref", source_id="cr1", title="A Study of AI Interview Assessment", doi="https://doi.org/10.123/abc", abstract="", source_database="crossref").to_dict(),
                ],
            )
            records, report = MetadataNormalizeDedupAgent().run(out)
            self.assertEqual(len(records), 1)
            self.assertEqual(report["duplicate_count"], 1)
            self.assertEqual(set(records[0].source_database), {"openalex", "crossref"})

    def test_abstract_only_record_is_not_deep_summarized_and_gets_request(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp)
            record = Record(
                record_id="r1",
                title="AI Literacy and Job Performance",
                abstract="This study examines AI literacy and job performance using survey data.",
                relevance_score=3,
                inclusion_status="included",
            )
            config = RunConfig(topic="AI literacy job performance", out=tmp)
            evidence = EvidenceAcquisitionAgent().run([record], config, out)
            summaries = DeepSummaryAgent().run([record], evidence, config, out)
            self.assertEqual(record.evidence_level, "abstract_only")
            self.assertEqual(summaries, [])
            self.assertEqual(record.summary_status, "needs_user_fulltext")
            requests = FullTextRequestAgent().run([record], evidence, config, out)
            self.assertEqual(len(requests), 1)
            self.assertEqual(requests[0]["record_id"], "r1")
            self.assertTrue((out / "fulltext_requests.jsonl").exists())
            self.assertTrue((out / "fulltext_requests.csv").exists())
            self.assertTrue((out / "fulltext_requests.md").exists())

    def test_evidence_acquisition_extracts_pdf_fulltext_for_deep_targets(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp)
            record = Record(
                record_id="r1",
                title="AI Literacy Full Text Study",
                abstract="This abstract is available.",
                open_access_pdf="https://example.test/paper.pdf",
                relevance_score=3,
                inclusion_status="included",
                citation_count=10,
            )
            pdf_text = """
Abstract
This paper studies AI literacy.

Introduction
AI literacy matters for work and learning.

Methods
Participants completed a survey and performance task.

Results
AI literacy predicted better task outcomes.

Conclusion
AI literacy should be developed through training.
"""
            with patch("litflow.agents.download_pdf", return_value=out / "fake.pdf"), patch("litflow.agents.extract_pdf_text", return_value=pdf_text):
                config = RunConfig(topic="AI literacy job performance", out=tmp, max_deep=1)
                evidence = EvidenceAcquisitionAgent().run([record], config, out)
            bundle = evidence["r1"]
            self.assertEqual(record.evidence_level, "fulltext_pdf")
            self.assertEqual(bundle["fulltext_status"], "completed")
            self.assertIn("Participants completed", bundle["fulltext_sections"]["method"])
            summaries = DeepSummaryAgent().run([record], evidence, config, out)
            self.assertIn("Participants completed", summaries[0]["method"])
            self.assertIn("AI literacy predicted", summaries[0]["results_findings"])

    def test_extract_elsevier_pii_from_crossref_links(self):
        self.assertEqual(
            extract_elsevier_pii("https://api.elsevier.com/content/article/PII:S2666920X24000262?httpAccept=text/xml"),
            "S2666920X24000262",
        )
        self.assertEqual(
            extract_elsevier_pii("https://www.sciencedirect.com/science/article/pii/S2666557324000247"),
            "S2666557324000247",
        )

    def test_packaging_creates_required_xlsx_sheets_and_docx(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp)
            write_json(out / "query_plan.json", {"queries": [{"source": "openalex", "query": "AI assessment", "intent": "test"}]})
            write_jsonl(out / "search_log.jsonl", [{"source": "fake", "query": "AI", "result_count": 1}])
            record = Record(
                record_id="r1",
                title="AI Assessment Validity",
                authors=["Kim"],
                year=2026,
                venue="Journal",
                abstract="Abstract",
                relevance_score=3,
                inclusion_status="included",
                evidence_level="abstract_only",
                source_database=["fake"],
            )
            summary = {
                "record_id": "r1",
                "citation": "Kim (2026). AI Assessment Validity.",
                "title": "AI Assessment Validity",
                "evidence_level": "abstract_only",
                "abstract": "Abstract",
                "introduction": "Intro",
                "method": "Method limited to abstract.",
                "results_findings": "Findings limited to abstract.",
                "conclusion": "Conclusion",
                "limitations": "Limitations",
                "topic_relevance": "Relevant",
                "follow_up": "Review full text.",
            }
            outputs = PackagingAgent().run([record], [summary], [], RunConfig(topic="AI assessment", out=tmp), out)
            wb = load_workbook(outputs["xlsx"], read_only=True)
            try:
                self.assertEqual(
                    set(["Run_Config", "Search_Log", "Queries", "All_Candidates", "Included", "Excluded", "Evidence_Level", "Deep_Summaries", "FullText_Requests", "QA_Flags"]),
                    set(wb.sheetnames),
                )
            finally:
                wb.close()
            with ZipFile(outputs["docx"]) as zf:
                xml = zf.read("word/document.xml").decode("utf-8", errors="ignore")
                self.assertIn("w:tblGrid", xml)
                self.assertNotIn("https://", xml)

    def test_mock_integration_run_creates_artifacts(self):
        with tempfile.TemporaryDirectory() as tmp:
            records = [
                RawRecord(
                    source="fake",
                    source_id="1",
                    title="Large Language Models for Interview Assessment",
                    authors=["Lee"],
                    year=2025,
                    venue="Conference",
                    doi="10.1000/test",
                    abstract="Large language models are used for interview assessment validity and reliability.",
                    url="https://example.test/paper",
                    citation_count=10,
                    source_database="fake",
                )
            ]
            connectors = {name: FakeConnector(name, records) for name in ["openalex", "crossref", "semantic_scholar", "arxiv", "web"]}
            out = Path(tmp) / "run"
            result = Orchestrator(connectors).run(RunConfig(topic="LLM interview assessment validity", out=str(out), max_raw=25, max_deep=5, agent_mode="off"))
            self.assertTrue(Path(result["outputs"]["xlsx"]).exists())
            self.assertTrue(Path(result["outputs"]["docx"]).exists())
            manifest = read_json(out / "run_manifest.json")
            self.assertEqual(manifest["stages"][-1]["name"], "packaging")
            included = [r for r in read_jsonl(out / "screened_records.jsonl") if r["inclusion_status"] == "included"]
            self.assertEqual(len(included), 1)
            self.assertTrue(included[0]["found_by"])
            self.assertGreaterEqual(included[0]["search_round"], 1)
            self.assertIn("facet_matches", included[0])

    def test_resume_fulltext_uses_existing_records_without_search(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "run"
            out.mkdir()
            record = Record(
                record_id="r1",
                title="AI Literacy Full Text Study",
                authors=["Kim"],
                year=2026,
                doi="10.123/fulltext",
                relevance_score=3,
                inclusion_status="included",
                evidence_level="abstract_only",
            )
            write_jsonl(out / "screened_records.jsonl", [record.to_dict()])
            write_json(out / "run_manifest.json", {"topic": "AI literacy", "config": {"topic": "AI literacy", "output_format": "files"}, "stages": []})
            user_dir = out / "user_fulltext"
            user_dir.mkdir()
            (user_dir / "10_123_fulltext.pdf").write_bytes(b"%PDF fake")
            with patch("litflow.agents.extract_pdf_text", return_value="Abstract\nFull text.\n\nMethods\nSurvey method.\n\nResults\nPositive result."):
                result = Orchestrator({"explode": ExplodingConnector()}).resume_fulltext(
                    RunConfig(topic="", out=str(out), output_format="files", agent_mode="off")
                )
            self.assertTrue(Path(result["outputs"]["xlsx"]).exists())
            summaries = read_jsonl(out / "summaries.jsonl")
            self.assertEqual(len(summaries), 1)
            requests = read_jsonl(out / "fulltext_requests.jsonl")
            self.assertEqual(requests, [])
            updated = read_jsonl(out / "screened_records.jsonl")[0]
            self.assertEqual(updated["evidence_level"], "user_provided_fulltext")
            self.assertEqual(updated["summary_status"], "completed")

    def test_sdk_assisted_agents_use_bridge_outputs(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp)
            bridge = FakeSDKBridge()
            config = RunConfig(topic="LLM interview assessment", out=tmp, agent_mode="sdk")
            scope = IntakeScopeAgent().run(config, out)
            query_plan = QueryStrategyAgent(bridge).run(scope, config, out)
            self.assertTrue(query_plan["sdk_assisted"])
            self.assertIn("LLM employee selection validity", [q["query"] for q in query_plan["queries"]])
            record = Record(
                record_id="r1",
                title="Large Language Models for Interview Assessment",
                abstract="LLM interview assessment validity.",
                open_access_pdf="https://example.test/sdk.pdf",
            )
            records = RelevanceScreeningAgent(bridge).run([record], scope, out)
            self.assertEqual(records[0].relevance_score, 3)
            with patch("litflow.agents.download_pdf", return_value=out / "sdk.pdf"), patch("litflow.agents.extract_pdf_text", return_value="Abstract\nSDK full text.\n\nMethods\nA method section."):
                evidence = EvidenceAcquisitionAgent().run(records, config, out)
            summaries = DeepSummaryAgent(bridge).run(records, evidence, config, out)
            self.assertEqual(summaries[0]["abstract"], "SDK abstract summary")
            flags = QualityAuditAgent(bridge).run(records, summaries, out, config.topic)
            self.assertEqual(flags[-1]["code"], "sdk_flag")

    def test_notion_page_id_parsing_and_rich_text_chunking(self):
        self.assertEqual(parse_notion_page_id("https://www.notion.so/Research-Hub-1234567890abcdef1234567890abcdef?pvs=4"), "1234567890abcdef1234567890abcdef")
        self.assertEqual(parse_notion_page_id("12345678-90ab-cdef-1234-567890abcdef"), "1234567890abcdef1234567890abcdef")
        chunks = rich_text("a" * 4100)
        self.assertEqual(len(chunks), 3)
        self.assertTrue(all(len(item["text"]["content"]) <= 1900 for item in chunks))

    def test_summary_record_to_notion_blocks_contains_sections(self):
        record = Record(record_id="r1", title="AI Assessment Validity", authors=["Kim"], year=2026, venue="Journal", evidence_level="abstract_only")
        summary = {
            "record_id": "r1",
            "citation": "Kim (2026). AI Assessment Validity.",
            "title": "AI Assessment Validity",
            "evidence_level": "abstract_only",
            "abstract": "Abstract summary",
            "introduction": "Intro summary",
            "method": "Method summary",
            "results_findings": "Findings summary",
            "conclusion": "Conclusion summary",
            "limitations": "Limitations summary",
            "topic_relevance": "Relevant",
            "follow_up": "Review full text",
        }
        blocks = summary_record_to_blocks(summary, record)
        block_types = [block["type"] for block in blocks]
        self.assertIn("callout", block_types)
        self.assertIn("divider", block_types)
        self.assertIn("heading_1", block_types)
        self.assertIn("heading_2", block_types)
        rendered = str(blocks)
        self.assertIn("Evidence Level", rendered)
        self.assertIn("Abstract summary", rendered)
        self.assertIn("Method summary", rendered)

    def test_hub_page_blocks_use_notion_native_design_structure(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp)
            write_json(out / "query_plan.json", {"queries": [{"source": "openalex", "query": "AI assessment", "intent": "test"}]})
            write_jsonl(out / "search_log.jsonl", [{"source": "openalex", "query": "AI", "result_count": 1}])
            write_json(out / "run_manifest.json", {"agents_sdk": {"enabled": False}})
            (out / "coverage_audit.md").write_text("Included records: 1\nQA flags: 0\n", encoding="utf-8")
            record = Record(record_id="r1", title="AI Assessment", relevance_score=3, inclusion_status="included", evidence_level="abstract_only")
            summary = {"record_id": "r1", "citation": "Kim (2026). AI Assessment.", "title": "AI Assessment"}
            blocks = hub_page_blocks(
                [record],
                [summary],
                [],
                RunConfig(topic="AI assessment", out=tmp),
                out,
                {"r1": {"url": "https://notion.test/paper", "page_id": "child1", "title": "Paper"}},
            )
            block_types = [block["type"] for block in blocks]
            self.assertIn("callout", block_types)
            self.assertIn("divider", block_types)
            self.assertLess(block_types.index("heading_1"), len(block_types))
            rendered = str(blocks)
            self.assertIn("Run Summary", rendered)
            self.assertIn("Full-Text Summaries", rendered)
            self.assertIn("Full Text Needed", rendered)
            self.assertIn("https://notion.test/paper", rendered)

    def test_notion_packaging_creates_and_updates_manifest_pages(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp)
            write_json(out / "query_plan.json", {"queries": [{"source": "openalex", "query": "AI assessment", "intent": "test"}]})
            write_jsonl(out / "search_log.jsonl", [{"source": "openalex", "query": "AI", "result_count": 1}])
            write_json(out / "run_manifest.json", {"agents_sdk": {"enabled": False}})
            (out / "coverage_audit.md").write_text("Included records: 1\nQA flags: 0\n", encoding="utf-8")
            record = Record(
                record_id="r1",
                title="AI Assessment Validity",
                authors=["Kim"],
                year=2026,
                venue="Journal",
                relevance_score=3,
                inclusion_status="included",
                evidence_level="abstract_only",
                source_database=["openalex"],
            )
            summary = {
                "record_id": "r1",
                "citation": "Kim (2026). AI Assessment Validity.",
                "title": "AI Assessment Validity",
                "evidence_level": "abstract_only",
                "abstract": "Abstract",
                "introduction": "Intro",
                "method": "Method",
                "results_findings": "Findings",
                "conclusion": "Conclusion",
                "limitations": "Limitations",
                "topic_relevance": "Relevant",
                "follow_up": "Review full text.",
            }
            fake = FakeNotionClient()
            config = RunConfig(topic="AI assessment", out=tmp, output_format="notion", notion_parent="1234567890abcdef1234567890abcdef")
            outputs = PackagingAgent(notion_client=fake).run([record], [summary], [], config, out)
            self.assertIn("notion", outputs)
            self.assertFalse((out / "literature_review.xlsx").exists())
            self.assertFalse((out / "literature_review.docx").exists())
            manifest = read_json(out / "notion_manifest.json")
            hub_id = manifest["hub_page_id"]
            child_id = manifest["child_pages"]["r1"]["page_id"]
            self.assertIn(hub_id, fake.pages_data)
            self.assertIn(child_id, fake.pages_data)

            summary["abstract"] = "Updated abstract"
            PackagingAgent(notion_client=fake).run([record], [summary], [], config, out)
            updated = read_json(out / "notion_manifest.json")
            self.assertEqual(updated["hub_page_id"], hub_id)
            self.assertEqual(updated["child_pages"]["r1"]["page_id"], child_id)
            self.assertIn(hub_id, fake.updated_pages)
            self.assertIn(child_id, fake.updated_pages)

    def test_notion_packaging_only_creates_child_pages_for_summaries(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp)
            write_json(out / "query_plan.json", {"queries": []})
            write_jsonl(out / "search_log.jsonl", [])
            write_json(out / "run_manifest.json", {"agents_sdk": {"enabled": False}})
            write_jsonl(
                out / "fulltext_requests.jsonl",
                [
                    {
                        "record_id": "missing",
                        "citation": "Lee (2026). Missing Full Text.",
                        "access_hint": "no_pdf_discovered",
                        "priority": 2,
                        "suggested_filename": "missing.pdf",
                    }
                ],
            )
            fulltext_record = Record(record_id="full", title="Full Text Paper", inclusion_status="included", evidence_level="fulltext_pdf")
            missing_record = Record(record_id="missing", title="Missing Full Text", inclusion_status="included", evidence_level="abstract_only", summary_status="needs_user_fulltext")
            summary = {
                "record_id": "full",
                "citation": "Kim (2026). Full Text Paper.",
                "title": "Full Text Paper",
                "evidence_level": "fulltext_pdf",
                "abstract": "Abstract",
                "research_purpose_questions": "Purpose",
                "theoretical_background": "Theory",
                "study_design_data_sample_context": "Design",
                "measures_variables_indicators": "Measures",
                "analysis_methods": "Analysis",
                "key_findings": "Findings",
                "discussion_contribution": "Discussion",
                "conclusion": "Conclusion",
                "limitations": "Limitations",
                "topic_relevance": "Relevant",
                "follow_up": "Follow up",
            }
            fake = FakeNotionClient()
            config = RunConfig(topic="AI assessment", out=tmp, output_format="notion", notion_parent="1234567890abcdef1234567890abcdef")
            PackagingAgent(notion_client=fake).run([fulltext_record, missing_record], [summary], [], config, out)
            manifest = read_json(out / "notion_manifest.json")
            self.assertEqual(set(manifest["child_pages"].keys()), {"full"})
            hub_blocks = str(fake.children[manifest["hub_page_id"]])
            self.assertIn("Full Text Needed", hub_blocks)
            self.assertIn("Missing Full Text", hub_blocks)

    def test_notion_packaging_requires_token_or_parent(self):
        with tempfile.TemporaryDirectory() as tmp:
            config = RunConfig(topic="AI assessment", out=tmp, output_format="notion", notion_token_env="LITFLOW_MISSING_TOKEN_FOR_TEST")
            with self.assertRaisesRegex(RuntimeError, "LITFLOW_MISSING_TOKEN_FOR_TEST"):
                PackagingAgent().run([], [], [], config, Path(tmp))

            config = RunConfig(topic="AI assessment", out=tmp, output_format="notion")
            with self.assertRaisesRegex(RuntimeError, "--notion-parent"):
                PackagingAgent(notion_client=FakeNotionClient()).run([], [], [], config, Path(tmp))


if __name__ == "__main__":
    unittest.main()
